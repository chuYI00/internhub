# -*- coding: utf-8 -*-
"""按岗位定制简历与投递材料（读取 data/profile.json，支持任何使用者换成自己的资料）。

用法：
  from ihub import resume
  resume.apply_message(job)              # 生成投递开场白
  resume.build_docx(job, 'out/cv.docx')  # 生成定制简历
资料在网页「👤 我的资料」里上传/修改，只存本机。
"""
import os
import re

from . import profile as profile_mod

# ---------------- 方向识别 ----------------
DIR_MEDIA = ["新媒体", "内容", "运营", "直播", "剪辑", "编辑", "文案", "编导", "市场",
             "传播", "短视频", "图文", "策划", "品牌", "营销", "主播"]
DIR_PM = ["产品经理", "产品", "PM", "需求", "商业化"]
DIR_AI = ["AI", "人工智能", "大模型", "智能体", "算法", "LangChain", "语音", "AIGC"]
DIR_HW = ["硬件", "嵌入式", "电子", "单片机", "MCU", "射频", "声学", "穿戴", "耳机", "音箱",
          "音频", "IoT", "物联网", "测试", "质量", "结构", "驱动", "固件", "民航", "机场"]
DIR_SOFT = ["开发", "后端", "前端", "软件", "数据", "运维", "IT", "测试开发", "工程师"]

INTENT = {
    "media": "新媒体 / 短视频内容运营 · AI 内容创作方向",
    "ai": "AI 应用 / 大模型方向",
    "pm": "产品经理（AI / 内容 / 硬件方向）",
    "hw": "嵌入式 / 物联网 / 硬件研发测试方向",
    "soft": "软件开发 / 测试 / IT 方向",
    "general": "综合方向（运营 / AI 应用 / 工程技术）",
}

# 使用者资料亮点不足时用于补齐的方向话术
DIR_FALLBACK = {
    "media": ["熟悉短视频内容生产全流程（选题-剪辑-发布-复盘）", "能用 AI 工具显著提升内容产出效率"],
    "ai": ["有 AI 应用开发实践（提示词工程 / 大模型调用）", "对 AI 工具上手快，能迁移到业务场景"],
    "pm": ["具备需求梳理与文档撰写能力", "能兼顾用户体验与技术可行性"],
    "hw": ["有软硬件项目动手经验，能读英文技术文档", "熟悉调试流程与问题排查"],
    "soft": ["掌握 Python / C 与 Linux 常用操作", "有完整项目从 0 到 1 的实践经验"],
    "general": ["学习能力强，能快速上手新工具", "沟通顺畅、执行到位"],
}


def detect_direction(job) -> str:
    blob = " ".join(str(job.get(k) or "") for k in
                    ("title", "company", "tags", "industry", "requirement")).lower()
    for keys, name in ((DIR_MEDIA, "media"), (DIR_AI, "ai"), (DIR_PM, "pm"),
                       (DIR_HW, "hw"), (DIR_SOFT, "soft")):
        if any(k in blob for k in keys):
            return name
    return "general"


def highlights_for(job, prof=None) -> list:
    prof = prof or profile_mod.load()
    hs = profile_mod.lines(prof.get("highlights"))
    d = detect_direction(job)
    if len(hs) >= 3:
        return hs[:3]
    return (hs + DIR_FALLBACK[d])[:3]


def intent_line(job) -> str:
    prof = profile_mod.load()
    t = str(job.get("title") or "该岗位").strip()
    c = str(job.get("company") or "").strip()
    city = str(job.get("city") or "").strip()
    target = f"{t}（{c}{('/' + city) if city and city != '全国' else ''}）"
    return f"应聘 {target}　|　{prof.get('intent') or INTENT[detect_direction(job)]}"


# 技能/要求词表（用于 JD 匹配与差距分析，可按需扩充）
VOCAB = [
    # 技术
    "物联网", "嵌入式", "单片机", "硬件", "电子", "通信", "计算机", "软件", "测试", "数据",
    "数据库", "算法", "人工智能", "大模型", "机器学习", "网络", "运维", "技术支持", "产品",
    "数据分析", "数据采集", "接口", "脚本", "提示词", "需求", "文档", "项目管理", "系统",
    "自动化", "电气", "控制", "云计算", "前端", "后端", "小程序", "安全", "民航", "机场",
    # 工具/语言
    "Python", "C++", "C语言", "Java", "SQL", "Excel", "PPT", "Linux", "Git", "MQTT", "Docker",
    "STM32", "OpenHarmony", "ROS", "SQLite", "Streamlit", "LangChain", "Figma", "PS", "剪映",
    # 通用能力
    "沟通", "协作", "抗压", "责任心", "学习能力", "执行", "逻辑", "细心", "服务意识",
    "新媒体", "内容", "视频", "剪辑", "设计", "运营", "市场", "销售", "客服", "财务", "人力", "行政",
]


def jd_analysis(jd_text: str, prof=None):
    """返回 (简历已覆盖的岗位要求, 岗位要求但简历未体现的)——用于精准贴合 + 诚实提示差距。"""
    prof = prof or profile_mod.load()
    jd = (jd_text or "").lower()
    if not jd.strip():
        return [], []
    mine = " ".join([
        str(prof.get("skills") or ""), str(prof.get("highlights") or ""),
        str(prof.get("experiences") or ""), str(prof.get("certificates") or ""),
        str(prof.get("major") or ""), str(prof.get("degree") or ""),
        str(prof.get("self_eval") or ""),
    ]).lower()
    # 词表命中 + JD 里的英文术语
    jd_terms = [v for v in VOCAB if v.lower() in jd]
    jd_terms += [t for t in re.findall(r"[A-Za-z][A-Za-z0-9+#./\-]{2,14}", jd_text or "")]
    seen, ordered = set(), []
    for t in jd_terms:
        if t.lower() not in seen:
            seen.add(t.lower()); ordered.append(t)
    hit = [t for t in ordered if t.lower() in mine]
    miss = [t for t in ordered if t not in hit]
    return hit[:20], miss[:20]


def _sort_by_jd(items, hit_keywords):
    """把与 JD 关键词重合度高的条目排前面（不新增内容，只重排，保证真实）。"""
    def score(s):
        return sum(1 for k in hit_keywords if k.lower() in str(s).lower())
    return sorted(items, key=score, reverse=True)


def apply_message(job) -> str:
    prof = profile_mod.load()
    t = str(job.get("title") or "该岗位").strip()
    name = prof.get("name") or "我"
    school = prof.get("school") or ""
    major = prof.get("major") or ""
    year = prof.get("graduate_year") or ""
    hl = highlights_for(job, prof)[0]
    return (f"您好！我是{school}{major}专业 {year} 届学生{name}，看到贵司正在招聘「{t}」，非常感兴趣并希望投递。"
            f"我的匹配点：{hl}。学习快、执行强，能稳定实习 2 个月以上。简历详见附件，期待与您进一步沟通！")


def safe_name(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s or "job")[:50]


# 兼容旧调用名
_simplify_path_safe = safe_name


def build_docx(job, out_path: str, jd_text: str = None) -> str:
    """生成按岗位定制的简历；jd_text 传入岗位 JD 时，会把简历里与该岗位相关的
    亮点/技能排到前面（只重排、不编造），让 HR 一眼看到最相关的信息。"""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    prof = profile_mod.load()
    hit_kws, _miss = ([], [])
    if jd_text:
        hit_kws, _miss = jd_analysis(jd_text, prof)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.2)
    sec.left_margin = sec.right_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def setf(r, size=10.5, bold=False, color="2A2A2A"):
        r.font.name = "微软雅黑"; r.font.size = Pt(size); r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        rPr = r._element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rPr.append(rf)
        rf.set(qn("w:eastAsia"), "微软雅黑")

    def para(txt="", size=10.5, bold=False, color="2A2A2A", align=None, after=3, before=0):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        setf(p.add_run(txt), size, bold, color)
        return p

    def heading(txt):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(3)
        setf(p.add_run(txt), 13, True, "015187")
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:space"), "2"); bot.set(qn("w:color"), "015187")
        b.append(bot); pPr.append(b)

    def bullet(label, text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.space_after = Pt(2)
        setf(p.add_run("▪ "), 10, False, "015187")
        if label:
            setf(p.add_run(label), 10.5, True)
        setf(p.add_run(text), 10.5)

    para(prof.get("name", ""), 22, True, "015187", WD_ALIGN_PARAGRAPH.CENTER, after=2)
    meta = "　|　".join([x for x in [prof.get("phone"), prof.get("email"), prof.get("city"),
                                     prof.get("gender")] if x])
    para(meta, 9, False, "6E6E6E", WD_ALIGN_PARAGRAPH.CENTER, after=1)
    edu = "　|　".join([x for x in [prof.get("school"), prof.get("major"),
                                   f"{prof.get('degree', '')}在读", prof.get("edu_range")] if x])
    para(edu, 9, False, "6E6E6E", WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(intent_line(job), 11, True, "015187", after=8)

    heading("教育背景")
    bullet("", edu)
    if prof.get("gpa"):
        bullet("GPA：", str(prof["gpa"]) + (f"，{prof['scholarship']}" if prof.get("scholarship") else ""))
    if prof.get("english"):
        bullet("英语：", prof["english"])

    heading("岗位匹配亮点")
    hl_list = highlights_for(job, prof)
    if hit_kws:
        hl_list = _sort_by_jd(hl_list, hit_kws)
    for h in hl_list:
        bullet("", h)

    heading("实践经历")
    exp_lines = profile_mod.lines(prof.get("experiences"))
    if hit_kws:
        exp_lines = _sort_by_jd(exp_lines, hit_kws)
    for line in exp_lines:
        title, rng, desc = profile_mod.parse_experience(line)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        setf(p.add_run(f"▎{title}"), 10.5, True)
        if rng:
            setf(p.add_run(f"　{rng}"), 9, False, "6E6E6E")
        if desc:
            pp = doc.add_paragraph(); pp.paragraph_format.left_indent = Cm(0.45)
            pp.paragraph_format.space_after = Pt(3)
            setf(pp.add_run(desc), 10)

    heading("技能与证书")
    skill_lines = profile_mod.lines(prof.get("skills"))
    if hit_kws:
        skill_lines = _sort_by_jd(skill_lines, hit_kws)
    for s in skill_lines:
        bullet("", s)
    if prof.get("certificates"):
        bullet("证书：", prof["certificates"])

    heading("自我评价")
    para(prof.get("self_eval", ""), 10.5, after=6)
    para("（本简历由 InternHub 按岗位自动整理，投递前请人工核对）", 8, False, "9AA5B1",
         WD_ALIGN_PARAGRAPH.CENTER)

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc.save(out_path)
    return out_path
