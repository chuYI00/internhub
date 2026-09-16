# -*- coding: utf-8 -*-
"""极简 Markdown → docx（够用于报告/说明：标题、段落、列表、引用、表格、粗体、行内代码）。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "微软雅黑"
ACCENT = "1F4E79"
DARK = "262626"
GREY = "6B6B6B"


def _font(run, size, bold=False, color=DARK, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), FONT if not mono else name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_rich(p, text, size=10.5, color=DARK):
    for seg in INLINE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            _font(p.add_run(seg[2:-2]), size, True, color)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            r = p.add_run(seg[1:-1])
            _font(r, size - 0.5, False, "9A3412", mono=True)
            r.font.highlight_color = None
        else:
            _font(p.add_run(seg), size, False, color)


def build(md_path, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = lines[i].strip()
                if not re.match(r"^\|[\s:\-|]+\|$", r):
                    rows.append([c.strip() for c in r.strip("|").split("|")])
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = r[ci] if ci < len(r) else ""
                    p = cells[ci].paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    add_rich(p, txt, 9.5, ACCENT if ri == 0 else DARK)
                    if ri == 0:
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if s.startswith("---") and set(s) <= set("-"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            b = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single")
            bot.set(qn("w:sz"), "6")
            bot.set(qn("w:space"), "1")
            bot.set(qn("w:color"), "D0D7DE")
            b.append(bot)
            pPr.append(b)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lv, txt = len(m.group(1)), m.group(2)
            size = {1: 18, 2: 14, 3: 12, 4: 11}[lv]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12 if lv <= 2 else 9)
            p.paragraph_format.space_after = Pt(4)
            if lv == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, txt, size, ACCENT)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(3)
            add_rich(p, s.lstrip("> ").strip(), 10, GREY)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(1.5)
            _font(p.add_run("• "), 10.5, False, ACCENT)
            add_rich(p, m.group(1), 10.5)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(1.5)
            add_rich(p, m.group(1) + ". " + m.group(2), 10.5)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        add_rich(p, s, 10.5)
        i += 1

    doc.save(out_path)
    return out_path
